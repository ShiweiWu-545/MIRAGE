
import os
import shutil
import sys
from math import nan

import pandas as pd

sys.path.append(os.path.dirname(os.path.dirname(__file__)))
import argparse
import torch

import pickle

sys.path.append('..')
from utils.misc import *
from utils.data_process import *
from torch.utils.data import Dataset, TensorDataset, ConcatDataset
import re
import logging
from logging import handlers
import openpyxl
import random
import datetime
import multiprocessing
import copy
from torch.utils.tensorboard import SummaryWriter
import time
from sklearn.model_selection import KFold
from Bio.PDB.Polypeptide import three_to_index, is_aa
import torch.optim as optim
from torch.utils.data import Dataset, ConcatDataset
from Bio.PDB import PDBParser
import glob
import statistics
from Bio import BiopythonWarning
import warnings
from models.load_ExteriorFeatures import *
from tqdm import tqdm
from utils.data_process import _mask_dict_recursively

import dataclasses
import io
import os

import h5py
from torch import cdist
from torch.utils.data import Dataset
import numpy as np
import pandas as pd

from pathlib import Path
from sklearn.metrics import pairwise_distances

from utils import proteins2
from utils.proteins2 import ProteinInput
import traceback
from Bio import PDB

ATOM_N, ATOM_CA, ATOM_C, ATOM_O, ATOM_CB = 0, 1, 2, 3, 4

one_to_three = {'A': 'ALA', 'R': 'ARG', 'N': 'ASN', 'D': 'ASP',
                'C': 'CYS', 'Q': 'GLN', 'E': 'GLU', 'G': 'GLY',
                'H': 'HIS', 'I': 'ILE', 'L': 'LEU', 'K': 'LYS',
                'M': 'MET', 'F': 'PHE', 'P': 'PRO', 'S': 'SER',
                'T': 'THR', 'W': 'TRP', 'Y': 'TYR', 'V': 'VAL',
                '-': '-'}

three_to_one = dict([val, key] for key, val in one_to_three.items())


def index_to_three(num):
    if num == 20:
        return '-'
    else:
        out = {}
        for v in one_to_three.values():
            if v == '-': continue
            out[three_to_index(v)] = v
        return out[num]








































































































































































































def get_PT_x_label(PT_path, reverseIsNewComplex=False):
    if isinstance(PT_path, str):
        datasets_PT = load_variable(PT_path)
    else:
        datasets_PT = PT_path
    compound_name_list = []
    x_list = []
    label_list = []
    for sample in datasets_PT:

        if sample[0]['pdb_name'].split('_')[0] == 'reverse':
            if reverseIsNewComplex:
                sample_name = sample[0]['pdb_name']
                sample_label = sample[0]['pdb_name'].split('_', 1)[0]
            else:
                sample_name = sample[0]['pdb_name'].split('_', 1)[1]
                sample_label = sample[0]['pdb_name'].split('_', 1)[0]
        else:
            sample_name = sample[0]['pdb_name']
            sample_label = ''

        if sample[0]['pdb_name'].split('_')[0] == 'HM':
            sample_compound = '_'.join(sample_name.split('_')[: 3])
        elif sample[0]['pdb_name'].split('_')[0] == 'reverse':
            sample_compound = '_'.join(sample_name.split('_')[: 3])
        else:
            sample_compound = '_'.join(sample_name.split('_')[: 2])
        x_list.append(sample_name)
        label_list.append(sample_label)
        compound_name_list.append(sample_compound)

    return x_list, label_list







































def pretraining_datasets(config, no_AbAg=False, S645=False, S4169=False, S8338=False, S83382=False, S1131=False,
                         M1707=False, S645_no27=False, validation_flag=False):
    train, validation, test = [], [], []
    datasets = ''
    if no_AbAg:
        train = load_variable(config.feature.choice_residue.join(config.datasets.train.no_AbAg.PT.split(' ')))
        test = load_variable(config.feature.choice_residue.join(config.datasets.test.S645.split(' ')))
    elif M1707:
        datasets = 'Skempi2'
        train = load_variable(config.feature.choice_residue.join(config.datasets.train.dataset_del_M1707.PT.split(' ')))
        test = load_variable(config.feature.choice_residue.join(config.datasets.test.M1707.split(' ')))

    train = load_ProtTans(train, datasets=datasets, clean=False)
    validation = load_ProtTans(validation, datasets=datasets, clean=False)
    test = load_ProtTans(test, datasets=datasets)

    if config.feature.datasets_extend_reverse:
        train = extend_reverse(config, train)
    return train, validation, test


def experiment_10fold_datasets_0(config, S1131=False, S645=False, S645_no27=False):
    py_name = config.model.fold_name
    if S1131:
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S1131.split(' ')))
    elif S645:
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645.split(' ')))
    elif S645_no27:
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' ')))


    file_num = py_name.split('_')[0][-1]
    train_data_list = [[], [], [], [], [], [], [], [], [], []]
    test_list = [[], [], [], [], [], [], [], [], [], []]
    kf = KFold(n_splits=10, shuffle=True, random_state=config.train.seed)
    num = 0
    for train_data_index, test_index in kf.split(data):

        for i in train_data_index:
            train_data_list[num].append(data[i])
        for j in test_index:
            test_list[num].append(data[j])
        num += 1
    train_data = train_data_list[int(file_num)]
    test = test_list[int(file_num)]


    train = []
    validation = []
    kf = KFold(n_splits=9, shuffle=True, random_state=config.train.seed)
    for train_index, validation_index in kf.split(train_data):
        for i in train_index:
            train.append(train_data[i])
        for j in validation_index:
            validation.append(train_data[j])
        break

    if config.feature.datasets_extend_reverse:
        train = extend_reverse(config, train)

    return train, validation, test


def experiment_10fold_datasets_1(config, fold=10, validation_flag=False, S8338=False, S83382=False, S4169=False,
                                 M1707=False, S1131=False, S645=False, S645_no27=False, ):
    py_name = config.model.fold_name
    name = ''
    datasets = ''
    if S1131:
        name = 'S1131'
        datasets = 'Skempi2'
        data_path = config.feature.choice_residue.join(config.datasets.test.S1131.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S1131.split(' ')))
    elif M1707:
        name = 'M1707'
        datasets = 'Skempi2'
        data_path = config.feature.choice_residue.join(config.datasets.test.M1707.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.M1707.split(' ')))
    elif S645_no27:
        name = 'S645_no27'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' ')))
    elif S645:
        name = 'S645'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645.split(' ')))
    elif S4169:
        name = 'S4169'
        datasets = 'Skempi2'
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))
    elif S8338:
        name = 'S4169'
        datasets = 'Skempi2'
        reverseIsNewComplex = True
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))
    elif S83382:
        name = 'S83382'
        datasets = 'Skempi2'
        name = 'S83382'
        reverseIsNewComplex = True
        data_path = config.feature.choice_residue.join(config.datasets.test.S8338.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S8338.split(' ')))

    else:
        print('error')
        return None

    train, validation, test = Intra_complex_division(config, name, data, fold, py_name, validation_flag=validation_flag,
                                                     intra=False)


    train = load_ProtTans(train, datasets=datasets, clean=False)
    validation = load_ProtTans(validation, datasets=datasets, clean=False)
    test = load_ProtTans(test, datasets=datasets)

    if S8338 == False:
        if config.feature.datasets_extend_reverse:
            train = extend_reverse(config, train)
    if S8338:
        train = extend_reverse(config, train)
        validation = extend_reverse(config, validation)
        test = extend_reverse(config, test)

    return train, validation, test


def get_ecod_rule_compound_dict(compound_dict):
    if not os.path.exists('../datasets/ecod_dict.ecod'):
        current_path = os.getcwd()
        cmd = 'python ' + current_path + '/../scripts/get_ECOD_Hname.py'
        os.system(cmd)
    ecod_dict = load_variable('../datasets/ecod_dict.ecod')


    out_dict = {}
    done_list = []
    for k, v in ecod_dict.items():
        if 'NO_H_NAME' in k:
            flag_name = False
        else:
            flag_name = True
        tem_list = []
        for i in v:
            try:
                compound_entries = compound_dict[i.upper()]
                done_list.append(i.upper())
                if flag_name:
                    tem_list += compound_entries
                else:
                    out_dict[i.upper()] = compound_entries
            except KeyError:
                continue
        if (tem_list) and flag_name:
            out_dict[k] = tem_list
    return out_dict, done_list


def get_rest_compound_dict(compound_dict, done_list):
    out_dict = {}
    all_list = list(compound_dict.keys())
    for one in all_list:
        if one in done_list:
            done_list.remove(one)
            continue
        out_dict[one] = compound_dict[one]

    return out_dict


def get_ecod_compound_dict(compound_dict):
    ecod_rule_compound_dict, done_list = get_ecod_rule_compound_dict(compound_dict)
    rest_compound_dict = get_rest_compound_dict(compound_dict, done_list)
    out_dict = {**ecod_rule_compound_dict, **rest_compound_dict}

    return out_dict


class DatasetsDivide(torch.utils.data.Dataset):
    def __init__(self, config, label, validation_flag):
        self.config = config
        self.datasets = label.split('_')[0]
        self.mission = label.split('_')[1]
        self.datasetsSplitClass = label.split('_')[2]
        self.validation_flag = validation_flag
        if self.mission == 'DG':
            set_dict = {
                'test1': ['Affinity_Benchmark_v5.5', 'PDBbind_v2020'],
                'test2': ['PDBbind_v2020'],
                'combined': ['Affinity_Benchmark_v5.5', 'PDBbind_v2020'],
                'test3': ['SKEMPI_v2.0'],
                'A1741': ['Affinity_Benchmark_v5.5', 'PDBbind_v2020']
            }


    def load_data(self, ):
        train_dataset, validation_dataset, test_dataset, contact_area_dict = self.mission_navigation()
        return train_dataset, validation_dataset, test_dataset, contact_area_dict

    def mission_navigation(self):
        if self.mission == 'DDG':
            train_dataset, validation_dataset, test_dataset, contact_area_dict = self.predict_DDG()
        elif self.mission == 'DG':
            train_dataset, validation_dataset, test_dataset, contact_area_dict = self.predict_DG()
        return train_dataset, validation_dataset, test_dataset, contact_area_dict

    def predict_DDG(self):
        data = self.datasets
        datasetsSplitClass = self.datasetsSplitClass
        validation_flag = self.validation_flag

        S645_flag, S645_no27_flag, S1131_flag, AF3S1131_flag, S4169_flag, S8338_flag, S8338_flag, S83382_flag, M1707_flag = False, False, False, False, False, False, False, False, False
        Ablation_flag = False

        if data == 'S645':
            S645_flag = True
        elif data == 'S645_no27':
            S645_no27_flag = True
        elif data == 'S1131':
            S1131_flag = True
        elif data == 'AF3S1131':
            AF3S1131_flag = True
        elif data == 'S4169':
            S4169_flag = True
        elif data == 'S8338':
            S8338_flag = True
        elif data == 'S83382':
            S83382_flag = True
        elif data == 'M1707':
            M1707_flag = True

        if 'Ablation' in data:
            Ablation_flag = True

        if datasetsSplitClass == 'structure':
            train_dataset, validation_dataset, test_dataset = experiment_fold_datasets_complex_level(config,
                                                                                                     validation_flag=validation_flag,
                                                                                                     S645_no27=S645_no27_flag,
                                                                                                     S645=S645_flag,
                                                                                                     S1131=S1131_flag,
                                                                                                     AF3S1131=AF3S1131_flag,
                                                                                                     M1707=M1707_flag,
                                                                                                     S4169=S4169_flag,
                                                                                                     S8338=S8338_flag,
                                                                                                     S83382=S83382_flag,
                                                                                                     )
        elif datasetsSplitClass == 'mut':
            train_dataset, validation_dataset, test_dataset = experiment_10fold_datasets_1(config,
                                                                                           validation_flag=validation_flag,
                                                                                           S645=S645_flag,
                                                                                           S1131=S1131_flag,
                                                                                           M1707=M1707_flag,
                                                                                           S4169=S4169_flag,
                                                                                           S8338=S8338_flag,
                                                                                           S83382=S83382_flag,
                                                                                           )
        elif datasetsSplitClass == 'structureLeaveOne':
            train_dataset, validation_dataset, test_dataset = experiment_leave_one_fold_datasets_complex_level(config,
                                                                                                               validation_flag=validation_flag,
                                                                                                               S645=S645_flag,
                                                                                                               S1131=S1131_flag,
                                                                                                               M1707=M1707_flag,
                                                                                                               S4169=S4169_flag,
                                                                                                               S8338=S8338_flag,
                                                                                                               S83382=S83382_flag,
                                                                                                               )
        elif datasetsSplitClass == 'blind':
            train_dataset, validation_dataset, test_dataset = pretraining_datasets(config,
                                                                                   validation_flag=validation_flag,
                                                                                   S645_no27=S645_no27_flag,
                                                                                   S645=S645_flag,
                                                                                   S1131=S1131_flag,
                                                                                   M1707=M1707_flag,
                                                                                   S4169=S4169_flag,
                                                                                   S8338=S8338_flag,
                                                                                   S83382=S83382_flag,
                                                                                   )

        contact_area_dict = load_variable(config.datasets.contact_area)
        return train_dataset, validation_dataset, test_dataset, contact_area_dict

    def load_blind_data(self, data):


        root_path = Path('../datasets')
        if data == 'S1131':
            set_path = root_path / (data + '.csv')
        else:
            set_path = root_path / (data + '.xlsx')

        if data == 'combined':
            test1_datasets = self.load_blind_data('test1')
            test2_datasets = self.load_blind_data('test2')
            combined_datasets = ConcatDataset([test1_datasets, test2_datasets])
            return combined_datasets
        elif data == 'PPB-Affinity':
            test1_datasets = self.load_blind_data('test1')
            test2_datasets = self.load_blind_data('test2')
            test3_datasets = self.load_blind_data('test3')
            testS1131_datasets = self.load_blind_data('S1131')
            testPPBdel_datasets = self.load_blind_data('PPB-Affinity_del_test123_S1131')
            PPB_datasets = ConcatDataset([test1_datasets, test2_datasets, test3_datasets, testS1131_datasets, testPPBdel_datasets])
            return PPB_datasets

        os.makedirs(root_path / self.config.feature.DGchoice_residue, exist_ok=True)
        if os.path.exists(root_path / self.config.feature.DGchoice_residue / (data + '_DG.pt')):
            data_datasets = load_variable(root_path / self.config.feature.DGchoice_residue / (data + '_DG.pt'))
            return data_datasets

        total_df = pd.read_excel(self.config.datasets_DG.DGdatasets)
        df = self.read_DGexcel(set_path, data)


        kwargs = {'Mutations': 'nan'}
        kwargs = {}
        matching_df, matching_index = self.find_matching_rows(df, total_df, **kwargs)


        data_datasets = load_PredictDG_datasets(self.config, matching_index, matching_df)
        save_variable(data_datasets, root_path / self.config.feature.DGchoice_residue / (data + '_DG.pt'))
        return data_datasets

    def read_DGexcel(self, set_path, data):
        try:
            df = pd.read_excel(set_path)
        except ValueError:
            df = pd.read_csv(set_path)
        if data == 'test3':
            def test3_pdb_column(pdb_value):
                return pdb_value.split('_')[0]

            df['PDB'] = df['#Pdb'].apply(test3_pdb_column)

            def test3_Pairwiseinteraction_column(Pairwiseinteraction_value):
                return ';'.join(Pairwiseinteraction_value.split('_')[1:])

            df['Pairwise interaction'] = df['#Pdb'].apply(test3_Pairwiseinteraction_column)

            def test3_Mutations_column(Mutations_value):
                if pd.isna(Mutations_value):
                    return np.nan
                Mutations_value = Mutations_value.split('_')
                out = ''
                for mut in Mutations_value:
                    out += mut[1] + '_' + mut[0] + mut[2:] + ', '

                return out.strip(' ,')

            df['Mutations'] = df['Mutation(s)_PDB'].apply(test3_Mutations_column)
        elif data == 'S1131':
            def S1131_pdb_column(pdb_value):
                return pdb_value.split('_')[0]

            df['PDB'] = df['#protein'].apply(S1131_pdb_column)

            def S1131_Pairwiseinteraction_column(Pairwiseinteraction_value):
                return ';'.join(Pairwiseinteraction_value.split('_'))

            df['Pairwise interaction'] = df['Partners(A_B)'].apply(S1131_Pairwiseinteraction_column)

            def S1131_Mutations_column(pdb_name, Mutations_value):
                return Mutations_value.replace(':', '_')









            df['Mutations'] = df.apply(lambda row: S1131_Mutations_column(row['#protein'], row['mutation']), axis=1)
            df = df.dropna(subset=['Mutations'])

            df_wt = df.drop_duplicates(subset='PDB', keep='first')
            df_wt['Mutations'] = np.nan
            df = pd.concat([df, df_wt])
        return df

    def blind_validation(self, config, train_data=None, test_data=None, validation_data=None):
        train_dataset = self.load_blind_data(train_data)
        test_dataset = self.load_blind_data(test_data)
        return train_dataset, '', test_dataset

    def predict_DG(self):
        data = self.datasets
        datasetsSplitClass = self.datasetsSplitClass
        validation_flag = self.validation_flag

        test1_flag, test2_flag, combined_flag, test3_flag, A1741_flag = False, False, False, False, False
        validation_data = ''
        train_data = ''
        if data == 'test1':
            test1_flag = True


            train_data = 'A1741'
            validation_data = ''
        elif data == 'test2':
            test2_flag = True
            train_data = 'A1741'

            validation_data = ''
        elif data == 'combined':
            train_data = 'A1741'

            validation_data = ''
            combined_flag = True
        elif data == 'test3':
            test3_flag = True
            train_data = 'A1741'
            validation_data = ''
        elif data == 'A1741':
            data = 'A1741'
            A1741_flag = True
        elif data == 'S1131':
            train_data = 'A1741'
            validation_data = ''

        if 'Ablation' in data:
            Ablation_flag = True

        if datasetsSplitClass == 'blind':
            train_dataset, validation_dataset, test_dataset = self.blind_validation(config,
                                                                                    train_data=train_data,
                                                                                    test_data=data,
                                                                                    validation_data=validation_data)
        elif datasetsSplitClass == 'random':
            train_dataset, validation_dataset, test_dataset = self.cross_validation(config,
                                                                                    data=data,
                                                                                    fold=5, )

        contact_area_dict = load_variable(config.datasets.contact_area)
        return train_dataset, validation_dataset, test_dataset, contact_area_dict

    def cross_validation(self, config, data, fold=10, validation_flag=False):
        py_name = config.model.fold_name

        data = self.load_blind_data(data)

        data_list = [[], [], [], [], [], [], [], [], [], []]
        kf = KFold(n_splits=fold, shuffle=True, random_state=config.train.seed)
        num = 0
        for train_data_index, test_index in kf.split(data):
            for j in test_index:
                data_list[num].append(data[j])
            num += 1


        file_num = py_name.split('_')[0][-1]
        for num in range(fold):
            print('subset', num, 'numbers:', len(data_list[num]))

        test = data_list[int(file_num)]
        validation = []
        random_num = None
        if validation_flag:
            num_list = [i for i in range(fold) if i != int(file_num)]
            random.seed(config.train.seed)
            random_num = num_list[random.randint(0, fold - 2)]
            validation = data_list[random_num]
        train = []
        for i in range(len(data_list)):
            if i == int(file_num): continue
            if i == random_num: continue
            train += data_list[i]

        random.seed(config.train.seed)
        random.shuffle(test)
        random.seed(config.train.seed)
        random.shuffle(train)
        random.seed(config.train.seed)
        random.shuffle(validation)

        return train, validation, test

    def find_matching_rows(self, df, total_df, **kwargs):

        results = []
        out_df = []
        out_list = []


        for index, row in df.iterrows():

            try:
                pdb_value = row['PDB']
            except KeyError:
                pdb_value = row['name']
            condition = (total_df['name'] == pdb_value.upper())

            condition &= (total_df['set'] != 'noready')
            debug0 = total_df[condition]


            try:
                Mutations = row['Mutations']
            except KeyError:
                Mutations = nan
            if not isinstance(Mutations, str):
                if np.isnan(Mutations):
                    condition &= (total_df['Mutations'].isna())
                else:
                    assert 'Mutations' == 0
            else:
                condition &= (total_df['Mutations'] == Mutations)
            debug_mut = total_df[condition]

            if self.datasets == 'test3':
                Mutations = row['Mutations']
                if not isinstance(Mutations, str):
                    if np.isnan(Mutations):
                        condition &= (total_df['Mutations'].isna())
                else:
                    condition &= (total_df['Mutations'] == Mutations)
                debug_mut = total_df[condition]


            if len(df) != 1741:
                try:
                    Pairwise_interaction = row['Pairwise interaction']
                except KeyError:
                    Pairwise_interaction = row['protein_pair'].strip('()').split("'")
                    Pairwise_interaction_0 = Pairwise_interaction[1].replace(', ', '')
                    Pairwise_interaction_1 = Pairwise_interaction[3].replace(', ', '')
                    Pairwise_interaction = Pairwise_interaction_0 + ';' + Pairwise_interaction_1

                proteinA = ', '.join(Pairwise_interaction.strip(';').split(';')[0])
                proteinB = ', '.join(Pairwise_interaction.strip(';').split(';')[1])

                condition &= (total_df['proteinA'].isin([proteinA.strip(' ,'), proteinB.strip(' ,')]))
                debug1 = total_df[condition]
                condition &= (total_df['proteinB'].isin([proteinA.strip(' ,'), proteinB.strip(' ,')]))
                debug2 = total_df[condition]

            for key, value in kwargs.items():
                if key == 'Mutations' and value == 'nan':
                    condition &= total_df['Mutations'].isna()
                    debug3 = total_df[condition]
                elif key == 'set' and isinstance(value, (list, tuple, str)):
                    if isinstance(value, str):
                        condition &= (total_df[key] == value)
                    else:
                        condition &= (total_df[key].isin(value))
                    debug4 = total_df[condition]
                elif key != 'Mutations':
                    condition &= (total_df[key] == value)
                    debug5 = total_df[condition]


            filtered_rows = total_df[condition]


            if not filtered_rows.empty:

                result_row = filtered_rows.iloc[-1]
                result_index = result_row.name
                results.append((result_row, result_index))
                out_df.append(result_row)
                out_list.append(result_index)
            else:

                if self.datasets == 'test3':
                    df_row = total_df.iloc[0]
                    new_row = pd.Series(index=df_row.index, dtype=df_row.dtypes)
                    new_row['set'] = 'SKEMPI_v2.0'
                    new_row['name'] = pdb_value
                    new_row['Mutations'] = get_mut_to_mutClean(pdb_value, Mutations)
                    new_row['proteinA'] = row['#Pdb'].split('_')[1]
                    new_row['proteinB'] = row['#Pdb'].split('_')[2]
                    new_row['dG'] = row['Binding_affinity']

                    out_df.append(new_row)
                    out_list.append('nan')
                print('lalala')
        out_df = pd.DataFrame(out_df)
        return out_df, out_list


def get_mut_to_mutClean(pdb_value, mut):
    result_dict = {}
    if os.path.exists('../debug/mut_to_mutClean.dict'):
        result_dict = load_variable('../debug/mut_to_mutClean.dict')
    else:
        skempi2 = pd.read_csv('../datasets/skempi_v2.csv', sep=';')
        for index, row in skempi2.iterrows():
            result_dict[row['#Pdb'].split('_')[0] + '_' + row['Mutation(s)_PDB']] = row['#Pdb'].split('_')[0] + '_' +\
                                                                                    row['Mutation(s)_cleaned']
        save_variable('../debug/mut_to_mutClean.dict')

    cleaned_mut = result_dict[pdb_value + '_' + mut_format_conversion(mut)]
    return cleaned_mut.split('_')[1]


def get_mutClean_to_mut(pdb_value, mut):
    result_dict = {}
    if os.path.exists('../debug/mutClean_to_mut.dict'):
        result_dict = load_variable('../debug/mutClean_to_mut.dict')
    else:
        skempi2 = pd.read_csv('../datasets/skempi_v2.csv', sep=';')
        for index, row in skempi2.iterrows():
            result_dict[row['#Pdb'].split('_')[0] + '_' + row['Mutation(s)_cleaned']] = row['#Pdb'].split('_')[
                                                                                            0] + '_' + row[
                                                                                            'Mutation(s)_PDB']
        save_variable(result_dict, '../debug/mutClean_to_mut.dict')
    try:
        mut = result_dict[pdb_value + '_' + mut_format_conversion(mut)]
    except KeyError:
        return None
    return mut.split('_')[1]


def divide_dataset(data_num, clusters, N):
    import numpy as np
    from collections import defaultdict

    cluster_sizes = {c: len(L) for c, L in clusters.items()}

    sorted_clusters = sorted(cluster_sizes, key=cluster_sizes.get, reverse=True)

    folds = defaultdict(list)

    AvgN = data_num / N

    used_clusters = set()
    foldid = 0

    for c in sorted_clusters:
        if len(folds) >= N:
            break
        folds[foldid].extend(clusters[c])
        used_clusters.add(c)
        foldid += 1


    threshold = AvgN + np.random.normal(0, 10)


    foldid = 0
    for c in sorted_clusters:
        if c in used_clusters:
            continue
        if len(folds[foldid]) + len(clusters[c]) < threshold:
            folds[foldid].extend(clusters[c])
            used_clusters.add(c)
        else:
            foldid = (foldid + 1) % N


    for c in sorted_clusters:
        if c not in used_clusters:
            folds[N - 1].extend(clusters[c])

    return folds


def write_data(compound_dict, data, fold, reverseIsNewComplex):

    data_list = []
    for i in range(fold):
        data_list.append([])
    for sample in data:

        if sample[0]['pdb_name'].split('_')[0] == 'reverse':
            if reverseIsNewComplex:
                sample_name = sample[0]['pdb_name']
                sample_label = sample[0]['pdb_name'].split('_', 1)[0]
            else:
                sample_name = sample[0]['pdb_name'].split('_', 1)[1]
                sample_label = sample[0]['pdb_name'].split('_', 1)[0]
        else:
            sample_name = sample[0]['pdb_name']
            sample_label = ''

        sample_index = [sample_name, sample_label]

        for num in range(fold):
            subset = compound_dict[num]
            if sample_index in subset:
                data_list[num].append(sample)
                break
    return data_list


def get_tvt_datasets(config, py_name, data_list, fold, validation_flag):

    file_num = py_name.split('_')[0][-1]
    for num in range(fold):
        print('subset', num, 'numbers:', len(data_list[num]))

    test = data_list[int(file_num)]
    validation = []
    random_num = None
    if validation_flag:
        num_list = [i for i in range(10) if i != int(file_num)]
        random.seed(config.train.seed)
        random_num = num_list[random.randint(0, 8)]
        validation = data_list[random_num]
    train = []
    for i in range(len(data_list)):
        if i == int(file_num): continue
        if i == random_num: continue
        train += data_list[i]

    random.seed(config.train.seed)
    random.shuffle(test)
    random.seed(config.train.seed)
    random.shuffle(train)
    random.seed(config.train.seed)
    random.shuffle(validation)

    print('train:', len(train))
    print('validation:', len(validation))
    print('test:', len(test))
    return train, validation, test


def get_chains(s):
    result = []
    previous_char = None

    for char in s:
        if char != previous_char:
            result.append(char)
            previous_char = char

    return result


def load_ProtTans(data, datasets, clean=False):
    ExteriorFeatures = ComplexLoader()

    warnings.filterwarnings("ignore")


    for x, y in tqdm(data):





        resseq = x['wt']['resseq']
        index = [x - 1 for x in resseq]
        chains = get_chains(x['wt']['chain_id'])
        assert len(chains) == len(set(chains))

        try:
            feature_dict = ExteriorFeatures.get_protein(x['pdb_name'], chains, datasets=datasets)

        except Exception as e:
            print(f"Error: {e}")

        mask = np.zeros(feature_dict['wt_prottrans'].shape[0]).astype('bool')
        mask[index] = True




        feature_dict['wt_prottrans'] = feature_dict['wt_prottrans'][mask]
        feature_dict['mut_prottrans'] = feature_dict['mut_prottrans'][mask]


        for k, v in feature_dict.items():
            label = k.split('_')[0]

            if label == 'wt':
                if isinstance(v, np.ndarray):
                    x['wt'][k.split('_')[-1]] = torch.from_numpy(v).to(x['wt']['pos14'])
            elif label == 'mut':
                if isinstance(v, np.ndarray):
                    x['mut'][k.split('_')[-1]] = torch.from_numpy(v).to(x['wt']['pos14'])
            else:
                x[k.split('_')[-1]] = v

    tem_num = []
    for num, i in enumerate(data):

        if clean:
            if i[1] == '8':
                tem_num.append(num)
    tem_num.sort(reverse=True)
    for i in tem_num:
        del data[i]

    return data


def experiment_fold_datasets_complex_level(config, ECOD=True, S645=False, S4169=False, S8338=False, S83382=False,
                                           S1131=False, AF3S1131=False, M1707=False, S645_no27=False, fold=10,
                                           validation_flag=False):
    py_name = config.model.fold_name
    name = ''
    datasets = ''
    reverseIsNewComplex = False
    if S1131:
        datasets = 'Skempi2'
        name = 'S1131'
        data_path = config.feature.choice_residue.join(config.datasets.test.S1131.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S1131.split(' ')))
    elif AF3S1131:
        datasets = 'Skempi2'

        name = 'S1131'
        data_path = config.feature.choice_residue.join(config.datasets.test.AF3S1131.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.AF3S1131.split(' ')))
    elif M1707:
        datasets = 'Skempi2'
        name = 'M1707'
        reverseIsNewComplex = True
        data_path = config.feature.choice_residue.join(config.datasets.test.M1707.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.M1707.split(' ')))
    elif S645_no27:
        datasets = 'ABbind'
        name = 'S645_no27'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' ')))
    elif S645:
        datasets = 'ABbind'
        name = 'S645'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645.split(' ')))
    elif S4169:
        datasets = 'Skempi2'
        name = 'S4169'
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))
    elif S8338:
        datasets = 'Skempi2'
        name = 'S4169'
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))


    elif S83382:
        datasets = 'Skempi2'
        name = 'S83382'
        reverseIsNewComplex = True
        data_path = config.feature.choice_residue.join(config.datasets.test.S8338.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S8338.split(' ')))
    else:
        print('error')
        return None




    data_list = []
    for i in range(fold):
        data_list.append([])


    x_list, label_list = get_PT_x_label(data_path, reverseIsNewComplex=reverseIsNewComplex)
    compound_dict = compound_classify(x_list, label_list)

    if ECOD:
        print('ECOD!!!')
        if os.path.exists('../datasets/' + name + '_' + str(fold) + 'fold_ECOD.pt'):
            print('load existing file!')
            data_list = load_variable('../datasets/' + name + '_' + str(fold) + 'fold_ECOD.pt')
            data_list = get_current_data_list(compound_dict, data, fold, data_list)
        else:
            print('Splitting the dataset by ECOD！')
            data_num = len(x_list)
            compound_dict = get_ecod_compound_dict(compound_dict)
            compound_dict = divide_dataset(data_num, compound_dict, fold)
            data_list = write_data(compound_dict, data, fold, reverseIsNewComplex=reverseIsNewComplex)
            save_variable(data_list, '../datasets/' + name + '_' + str(fold) + 'fold_ECOD.pt')
        train, validation, test = get_tvt_datasets(config, py_name, data_list, fold, validation_flag)

        train = load_ProtTans(train, datasets=datasets, clean=False)
        validation = load_ProtTans(validation, datasets=datasets, clean=False)
        test = load_ProtTans(test, datasets=datasets)

        if S8338 == False:
            if config.feature.datasets_extend_reverse:
                train = extend_reverse(config, train)
        if S8338:
            train = extend_reverse(config, train)
            validation = extend_reverse(config, validation)
            test = extend_reverse(config, test)






        return train, validation, test


def get_current_data_list(compound_dict, data, fold, reference_data_list):

    data_list = []
    for i in range(fold):
        data_list.append([])


    reference_dict = {}
    for num, subset in enumerate(reference_data_list):
        for i in subset:
            tem_pdb_complex_name = i[0]['pdb_name']
            if tem_pdb_complex_name.split('_')[0] == 'reverse':
                pdb_complex_name = '_'.join(tem_pdb_complex_name.split('_')[0: 3])
            else:
                pdb_complex_name = '_'.join(tem_pdb_complex_name.split('_')[0: 2])
            try:
                reference_dict[pdb_complex_name].append(num)
            except KeyError:
                reference_dict[pdb_complex_name] = [num]
    for i in reference_dict.keys():
        reference_dict[i] = set(reference_dict[i])


    for sample in data:

        if sample[0]['pdb_name'].split('_')[0] == 'reverse':
            pdb_complex_name = '_'.join(sample[0]['pdb_name'].split('_')[0: 3])
        else:
            pdb_complex_name = '_'.join(sample[0]['pdb_name'].split('_')[0: 2])

        subset_num = reference_dict[pdb_complex_name]
        assert len(subset_num) == 1
        subset_num = list(subset_num)[0]
        data_list[subset_num].append(sample)

    return data_list


def get_new_f(dir_path, ):
    f_list = []
    for curDir, dirs, files in os.walk(dir_path):
        f_list = files
        break
    f_path_list = []
    f_existence_time = []
    f_num = []
    for f in f_list:
        f_path = dir_path + '/' + f
        if f.split('_')[0] == 'train':
            f_num.append(int(f.split('_')[-1].split('.')[0]))
        else:
            f_num.append(int(f.split('_')[0]))
        f_path_list.append(f_path)
        f_existence_time.append(get_file_existence_time(f_path))



    if len(set(f_existence_time)) == 1:
        min_existence_time_index = f_num.index(max(f_num))
    else:
        min_existence_time_index = f_existence_time.index(min(f_existence_time))
    f = os.path.basename(f_path_list[min_existence_time_index])

    return f_path_list[min_existence_time_index], f


def load_default_settings(config, py_name, label, validation_flag, debug=True, load_model='validation_model',
                          predict_flag=False, continue_train=False):

    if predict_flag:
        validation_flag = False
        debug = False
    weight = None
    weight_name = None


    if label.split('_')[-1] == 'structureLeaveOne':
        config = get_complex_tested(config, data='S645')
        py_name = config.model.fold_name
    else:
        config.model.fold_name = py_name



    set_seed(config.train.seed)
    logger = log_test(config.model.fold_name.split(',')[0])


    if continue_train:
        print('Continue train!!!')

        for curDir, dirs, files in os.walk('../data/mymodel/' + py_name.split('.')[0] + '/'):
            break
        for i in files:
            if 'best' in i and i.split('.')[-1] == 'pt':
                break

        weight, config = load_variable('../data/mymodel/' + py_name.split('.')[0] + '/' + i)
        config.train.max_iters = 10000
        config.feature.energy = False

        return weight, config, weight_name, py_name, logger

    if validation_flag:
        config.observer.validation_Rp = True
        config.observer.validation_losses_epoch = True

    else:
        if debug == False:
            weight, config = load_result_model(load_model, predict_flag, label, py_name)
        else:

            config.observer.validation_Rp = False
            config.observer.validation_losses_epoch = False





    return weight, config, weight_name, py_name, logger


def load_result_model(load_model, predict_flag, label, py_name):
    if load_model == 'validation_model':
        print('load validation config!!!')
        model_config_path = '../result/' + label + '/validation_result/' + py_name.split('.')[0] + '/'
    elif load_model == 'test_model':
        print('load test config!!!')
        model_config_path = '../result/' + label + '/' + py_name.split('.')[0] + '/'
    elif load_model == 'best_model':
        print('load best_model config!!!')
        model_config_path = '../result/' + label + '/'
    elif load_model == 'latest_model':
        print('load latest_model config!!!')


    for curDir, dirs, files in os.walk(model_config_path):
        break
    f = 'error'
    for f in files:

        if f.split('.')[-1] == 'pt':
            break
    assert f != 'error'
    model_config_path = model_config_path + f
    weight, config = load_variable(model_config_path)

    config.observer.validation_Rp = False
    config.observer.validation_losses_epoch = False
    config.train.max_iters = int(f.split('.')[0].split('_')[-1])

    config.train.max_iters = 200
    config.feature.energy = False


    config.feature.weight.Attention = False
    config.feature.Side_chain_geometry = False
    return weight, config


def get_batch(train_dataset, validation_dataset, test_dataset, config):
    train_dataset = batch_generator(train_dataset, batch_size=config.train.batch_size)
    validation_dataset = batch_generator(validation_dataset, batch_size=config.train.batch_size)
    test_dataset = batch_generator(test_dataset, batch_size=config.train.batch_size)
    return train_dataset, validation_dataset, test_dataset


def load_model_data(config, label, validation_flag=True):
    DatasestDivide = DatasetsDivide(config, label, validation_flag)
    train_dataset, validation_dataset, test_dataset, contact_area_dict = DatasestDivide.load_data()
    return train_dataset, validation_dataset, test_dataset, contact_area_dict


def Intra_complex_division(config, name, data, fold, py_name, validation_flag=False, intra=True):
    if os.path.exists('../datasets/' + name + '_' + str(fold) + 'fold_MUT.pt'):
        print('load existing file!')
        data_list = load_variable('../datasets/' + name + '_' + str(fold) + 'fold_MUT.pt')
    else:
        if not intra:
            data_list = [[], [], [], [], [], [], [], [], [], []]
            kf = KFold(n_splits=10, shuffle=True, random_state=config.train.seed)
            num = 0
            for train_data_index, test_index in kf.split(data):
                for j in test_index:
                    data_list[num].append(data[j])
                num += 1
        else:

            data_list = [[] for _ in range(fold)]


            x_list, label_list = get_PT_x_label(data)
            compound_dict = compound_classify(x_list, label_list)


            data_name_list = [[] for _ in range(fold)]
            for num, (k, v) in enumerate(compound_dict.items()):
                random.seed(config.train.seed)
                random.shuffle(v)
                group_subsets = [[] for _ in range(fold)]

                for i, sample in enumerate(v):
                    group_subsets[i % 10].append(sample)

                subset_index_list = list(range(10))
                random.seed(num)
                random.shuffle(subset_index_list)
                for i in range(fold):
                    data_name_list[i].extend(group_subsets[subset_index_list[i]])


            for sample in data:

                if sample[0]['pdb_name'].split('_')[0] == 'reverse':
                    sample_name = sample[0]['pdb_name'].split('_', 1)[1]
                    sample_label = sample[0]['pdb_name'].split('_', 1)[0]
                else:
                    sample_name = sample[0]['pdb_name']
                    sample_label = ''
                only_index = [sample_name, sample_label]
                for num in range(fold):
                    subset = data_name_list[num]
                    if only_index in subset:
                        data_list[num].append(sample)
                        break
        save_variable(data_list, '../datasets/' + name + '_' + str(fold) + 'fold_MUT.pt')


    file_num = py_name.split('_')[0][-1]
    for num in range(fold):
        print('subset', num, 'numbers:', len(data_list[num]))

    test = data_list[int(file_num)]
    validation = []
    random_num = None
    if validation_flag:
        num_list = [i for i in range(10) if i != int(file_num)]
        random.seed(config.train.seed)
        random_num = num_list[random.randint(0, 8)]
        validation = data_list[random_num]
    train = []
    for i in range(len(data_list)):
        if i == int(file_num): continue
        if i == random_num: continue
        train += data_list[i]

    random.seed(config.train.seed)
    random.shuffle(test)
    random.seed(config.train.seed)
    random.shuffle(train)
    random.seed(config.train.seed)
    random.shuffle(validation)

    return train, validation, test


def get_model_epoch(config, validation_flag):
    py_name = config.model.fold_name

    if validation_flag is False:
        py_name_list = load_variable('../debug/validation_best_model_index/model_list.pt')
        model_list = load_variable('../debug/validation_best_model_index/max_index_list.pt')

        for name, train_best_model in zip(py_name_list, model_list):
            if name == py_name.split('.')[0]:
                model_epoch = int(train_best_model.split('.')[0].split('_')[-1])
                break
        config.train.max_iters = model_epoch
    return config


def experiment_leave_one_fold_datasets_complex_level(config, validation_flag=False, S4169=False, S83382=False,
                                                     S8338=False, S645=False, S1131=False, M1707=False, S645_no27=False,
                                                     fold=10):
    py_name = config.model.fold_name
    if S1131:
        datasets = 'Skempi2'
        name = 'S1131'
        data_path = config.feature.choice_residue.join(config.datasets.test.S1131.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S1131.split(' ')))
    elif M1707:
        datasets = 'Skempi2'
        name = 'M1707'
        data_path = config.feature.choice_residue.join(config.datasets.test.M1707.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.M1707.split(' ')))
    elif S645_no27:
        datasets = 'ABbind'
        name = 'S645_no27'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' ')))
    elif S645:
        datasets = 'ABbind'
        name = 'S645'
        data_path = config.feature.choice_residue.join(config.datasets.test.S645.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645.split(' ')))
    elif S4169:
        datasets = 'Skempi2'
        name = 'S4169'
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))
    elif S8338:
        datasets = 'Skempi2'
        name = 'S4169'
        data_path = config.feature.choice_residue.join(config.datasets.test.S4169.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S4169.split(' ')))


    elif S83382:
        datasets = 'Skempi2'
        name = 'S83382'
        reverseIsNewComplex = True
        data_path = config.feature.choice_residue.join(config.datasets.test.S8338.split(' '))
        data = load_variable(config.feature.choice_residue.join(config.datasets.test.S8338.split(' ')))
    else:
        print('error')
        return None


    x_list, label_list = get_PT_x_label(data_path)
    compound_dict = compound_classify(x_list, label_list)
    compound_list = list(compound_dict.keys())


    data_dict = {}
    for i in compound_list:
        data_dict[i] = []















    for sample in data:

        if sample[0]['pdb_name'].split('_')[0] == 'reverse':
            sample_name = sample[0]['pdb_name']
            sample_label = sample[0]['pdb_name'].split('_', 1)[0]
        else:
            sample_name = sample[0]['pdb_name']
            sample_label = ''


        if sample_name.split('_')[0] == 'HM':
            sample_compound = sample_name.split('_')[1]
        elif sample_name.split('_')[0] == 'reverse':
            sample_compound = '_'.join(sample_name.split('_')[: 2])
        else:
            sample_compound = sample_name.split('_')[0]


        data_dict[sample_compound].append(sample)


    file_num = py_name.rsplit('/')[-1].rsplit('_', 1)[0]
    test = data_dict[file_num]
    train = []
    for i in compound_list:
        if i == file_num: continue
        train += data_dict[i]

    random.seed(config.train.seed)
    random.shuffle(test)
    random.seed(config.train.seed)
    random.shuffle(train)

    if config.feature.datasets_extend_reverse:
        train = extend_reverse(config, train)

    return train, None, test


def get_complex_tested(config, data='S1131'):
    S645_flag, S645_no27_flag, S1131_flag, S4169_flag, S8338_flag, S8338_flag, S83382_flag, M1707_flag = False, False, False, False, False, False, False, False

    if data == 'S645':
        S645_flag = True
    elif data == 'S645_no27':
        S645_no27_flag = True
    elif data == 'S1131':
        S1131_flag = True
    elif data == 'S4169':
        S4169_flag = True
    elif data == 'S8338':
        S8338_flag = True
    elif data == 'S83382':
        S83382_flag = True
    elif data == 'M1707':
        M1707_flag = True


    if S1131_flag:
        data_path = config.feature.choice_residue.join(config.datasets.test.S1131.split(' '))

    elif M1707_flag:
        data_path = config.feature.choice_residue.join(config.datasets.test.M1707.split(' '))

    elif S645_no27_flag:
        data_path = config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' '))

    elif S645_flag:
        data_path = config.feature.choice_residue.join(config.datasets.test.S645.split(' '))

    else:
        print('error')
        return None


    x_list, label_list = get_PT_x_label(data_path)
    compound_dict = compound_classify(x_list, label_list)
    compound_list = list(compound_dict.keys())


    if not os.path.exists('./compound_list.csv'):
        with open('./compound_list.csv', 'w') as f:
            f.write(str(0) + ':Index of the next complex to be tested.' + '\n')
            for i in compound_list:
                f.write(i + '\n')

    complex_data = open('./compound_list.csv', 'r').readlines()
    out_f = open('./compound_list.csv', 'w')
    index_num = complex_data[0].split(':')[0].strip()
    complex_list = complex_data[1:]


    out_f.write(str(int(index_num) + 1) + ':Index of the next complex to be tested.' + '\n')
    for i in complex_list:
        out_f.write(i)
    out_f.close()

    complex_name = complex_list[int(index_num)].strip()
    py_name = 'fold_one_complex_model/' + complex_name + '_model'
    config.model.fold_name = py_name
    return config












































def load_assign_datasets(config, train='', validation='', test=''):
    train_data, validation_data, test_data = [], [], []
    if train == 'S645_no27':
        train_data = load_variable(config.feature.choice_residue.join(config.datasets.test.S645_no27.split(' ')))

    if test == 'HIV_78':
        test_data = load_variable(config.feature.choice_residue.join(config.datasets.test.HIV_78.split(' ')))

    if validation == 'AB_228':
        validation_data = load_variable(config.feature.choice_residue.join(config.datasets.test.AB_228.split(' ')))

    train_data = extend_reverse(config, train_data)
    return train_data, validation_data, test_data




































def compound_classify(x_list, label_list):
    compound_dict = {}
    for x, label in zip(x_list, label_list):
        name = x.split('_')
        if name[0] == 'HM':

            compound = name[1]
        elif name[0] == 'reverse':
            compound = name[0] + '_' + name[1]
        else:

            compound = name[0]
        if compound in compound_dict:
            muts_list = compound_dict[compound]
        else:
            muts_list = []
            compound_dict[compound] = muts_list
        muts_list.append([x, label])
    return compound_dict


def batch_generator(data, batch_size):
    if not data: return None


    split_data = []

    num = 0
    tem = []
    for i in data:



        num += 1
        tem.append(i)
        if num == batch_size:
            split_data.append(tem)
            num = 0
            tem = []
            continue
    if num != 0:
        split_data.append(tem)

    collate_fn = PaddingCollate()
    out = []
    for chunks in split_data:
        x_list = []
        y_list = []
        for i in chunks:
            x_list.append(i[0])
            y_list.append(i[1])
        x_batch = collate_fn(x_list)
        y_batch = y_list
        out.append((x_batch, y_batch))
    return out


def log_test(dir):

    if not os.path.exists('./log/' + dir):
        os.makedirs('./log/' + dir)

    logger = logging.getLogger('debug')


    logger.setLevel(level=logging.DEBUG)


    formatter = logging.Formatter('%(asctime)s - %(filename)s[line:%(lineno)d] - %(levelname)s: %(message)s')


    stream_handler = logging.StreamHandler()
    stream_handler.setLevel(logging.DEBUG)
    stream_handler.setFormatter(formatter)
    logger.addHandler(stream_handler)

    time_rotating_file_handler = handlers.TimedRotatingFileHandler(filename='./log/' + dir + '/time_rotating_test.log',
                                                                   when='H')
    time_rotating_file_handler.setLevel(logging.INFO)
    time_rotating_file_handler.setFormatter(formatter)
    logger.addHandler(time_rotating_file_handler)

    return logger


def tensorboard_log(py_name):

    if not os.path.exists("../scripts/tensorboard_log"):
        os.makedirs("../scripts/tensorboard_log")
    timestamp = time.strftime("/%Y%m%d-%H%M%S" + "/")
    writer = SummaryWriter("./tensorboard_log/" + py_name + timestamp)
    return writer


def extend_reverse(config, data):
    out = []
    for sample in data:
        out.append(sample)
        reverse_sample = list(copy.deepcopy(sample))
        for i in reverse_sample:
            if isinstance(i, dict):
                tem = copy.deepcopy(i)
                i['wt'] = tem['mut']
                i['mut'] = tem['wt']


                i['pdb_name'] = 'reverse_' + tem['pdb_name']
            if isinstance(i, str):
                tem = copy.deepcopy(i)
                reverse_sample[1] = y_reverse(tem)
        out.append(tuple(reverse_sample))
    random.seed(config.train.seed)
    random.shuffle(out)
    return out


def one_hot(x, bins):
    if x.size()[-1] != 1:
        x = x.unsqueeze(-1)
    if len(x.size()) == 3:
        i, j, k = x.size()
        p = torch.full((i, j, len(bins)), fill_value=0).to(x)
        b = torch.argmin(torch.abs(x - (torch.tensor(bins).to(x))[None, None, :]), dim=2)
        p = p.scatter_(2, b.unsqueeze(2), 1)

    elif len(x.size()) == 4:
        i, j, k, u = x.size()
        p = torch.full((i, j, k, len(bins)), fill_value=0).to(x)
        b = torch.argmin(torch.abs(x - (torch.tensor(bins).to(x))[None, None, :]), dim=3)
        p = p.scatter_(3, b.unsqueeze(3), 1)

    return p


def outer_sum(x, y):
    pass


def ten_max_min_mean_var(tensor, name=''):
    max_x = torch.max(tensor)
    min_x = torch.min(tensor)
    mean_x = torch.mean(tensor)
    var_x = torch.var(tensor)
    print(name + 'max:', max_x)
    print(name + 'min:', min_x)
    print(name + 'mean:', mean_x)
    print(name + 'var:', var_x)


class CustomDataset(Dataset):

    def __init__(self, config, x_data_available=None, x_data=None, y_data=None, S1131=False, M1707=False, total=False,
                 S645=False, S645_no27=False, C3684=False, predict=False, no_AbAg=False, S4169=False, S8338=False):
        self.config = config
        if isinstance(x_data, list) and isinstance(y_data, list):
            random.seed(self.config.train.seed)
            random.shuffle(x_data)
            random.seed(self.config.train.seed)
            random.shuffle(y_data)
            self.x, self.y = x_data, y_data
        else:
            self.x, self.y = self.Load_Datasets(x_data_available, x_data, y_data, S1131, M1707, total, S645, S645_no27,
                                                C3684, predict, no_AbAg, S4169, S8338)

    def __getitem__(self, index):




        return self.x[index], self.y[index]

    def __len__(self):
        return len(self.x)

    @classmethod
    def pool_process(cls, n=os.cpu_count()):
        manager = multiprocessing.Manager()
        x = manager.list([])
        y = manager.list([])
        error = manager.list([])
        lock = manager.Lock()
        lock_file = manager.Lock()
        lock_file_csv = manager.Lock()
        po = multiprocessing.Pool(n)
        return po, x, y, error, lock, lock_file, lock_file_csv

    def load_data(self, x_available_list, x_list, y_dict, x_data_available, label_list=None):

        if label_list == None:
            label_list = [''] * len(x_list)

        pool = self.config.model.pool
        if pool:
            po, x, y, error, lock, lock_file, lock_file_csv = self.pool_process()
        else:
            x = []
            y = []
            error = []

        x_path_dict = {}

        for filer in x_available_list:
            pdb_path = x_data_available + '/' + filer
            pdb_tem = [0, 0]
            for pdb in os.listdir(pdb_path):
                if 'MUT_' in pdb and pdb.split('.')[-1] == 'pdb':
                    pdb_tem[0] = x_data_available + '/' + filer + '/' + pdb
                elif 'WT_' in pdb and pdb.split('.')[-1] == 'pdb':
                    pdb_tem[1] = x_data_available + '/' + filer + '/' + pdb
            x_path_dict[filer] = pdb_tem








        file_IndexError, file_no_y, file_duplication_y, file_no_Cddg, file_mut_not_on_surface = self.error_f()


        df = pd.DataFrame()
        df['core_ratio'] = None
        df['distance(core center-mut)'] = None
        df.to_csv('../error_record/core_ratio.csv', index=False)

        num = 0
        for pdb, label in zip(x_list, label_list):
            num += 1










            finished_num = str(num) + '/' + str(len(x_list))
            print('--------------------正在载入第%d个pdb---------------' % num)

            error_flag = False
            try:
                print(y_dict[pdb])
            except KeyError:
                file_no_y.write(label + pdb + '\n')
                error_flag = True



            if not error_flag:
                if y_dict[pdb] == [0]:
                    file_duplication_y.write(label + pdb + '\n')
                    error_flag = True
            try:
                print(x_path_dict[pdb])
            except KeyError:
                file_no_Cddg.write(label + pdb + '\n')
                error_flag = True
            if error_flag: continue






            if not pool:
                x, y, error = load_wt_mut_pdb_pair(x, y, error,
                                                   x_path_dict[pdb][1], x_path_dict[pdb][0],
                                                   pdb, label,
                                                   y_dict[pdb],
                                                   finished_num)
            else:









                po.apply_async(load_wt_mut_pdb_pair, (x, y, error,
                                                      x_path_dict[pdb][1], x_path_dict[pdb][0],
                                                      pdb, label,
                                                      y_dict[pdb],
                                                      finished_num,
                                                      lock, lock_file, lock_file_csv))


        if pool:
            po.close()
            po.join()


        for e in list(error):
            file_IndexError.write(e + '\n')
        file_IndexError.close()
        file_no_Cddg.close()
        file_no_y.close()
        file_duplication_y.close()
        file_mut_not_on_surface.close()
        x = list(x)
        y = list(y)
        return x, y

    def error_f(self):

        file_IndexError = open('../error_record/no_mutation_point_error.csv', 'w')
        file_no_y = open('../error_record/no_y.csv', 'w')
        file_duplication_y = open('../error_record/duplication_y.csv', 'w')
        file_no_Cddg = open('../error_record/no_Cddg.csv', 'w')
        file_mut_not_on_surface = open('../error_record/mut_not_on_surface.csv', 'w')
        return file_IndexError, file_no_y, file_duplication_y, file_no_Cddg, file_mut_not_on_surface

    def surface_filter(self, sample, error_flag):
        if error_flag: return '_', error_flag

        mut_whether_surface = []
        return error_flag, mut_whether_surface

    @classmethod
    def pdb_neme_toAB(cls, i):
        mut = None
        if isinstance(i, list):
            mut = i[2]
            i = i[0]
        tem = i.split('_', 1)
        tem[1] = tem[1].replace('_', '')
        tem = tem[0] + '_' + tem[1]

        if mut != None:
            mut = mut.replace(',', '_')
            pdb = tem + '_' + mut
        else:
            pdb = tem
        return pdb

    def Load_Datasets(self, x_data_available, x_data, y_data, S1131, M1707, total, S645, S645_no27, C3684, predict,
                      no_AbAg, S4169, S8338):

        if S1131:
            x, y = self.Load_S1131_Datasets(x_data_available, x_data, y_data, S1131)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif S4169:
            x, y = self.Load_S4169_Datasets(x_data_available, x_data, y_data, S4169)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif S8338:
            x, y = self.Load_S8338_Datasets(x_data_available, x_data, y_data, S8338)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif M1707:
            x, y = self.Load_M1707_Datasets(x_data_available, x_data, y_data, M1707)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif total:
            x, y = self.Load_total_Datasets(x_data_available, x_data, y_data, S1131, M1707, total)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif S645:
            x, y = self.Load_S645_Datasets(x_data_available, x_data, y_data, S645)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif S645_no27:
            x, y = self.Load_S645_no27_Datasets(x_data_available, x_data, y_data, S645_no27)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif predict:
            x, y = self.Load_predict_Datasets(x_data_available, x_data, y_data, predict)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif C3684:
            x, y = self.Load_C3684_Datasets(x_data_available, x_data, y_data, C3684)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y
        elif no_AbAg:
            x, y = self.Load_no_AbAg_Datasets(x_data_available, x_data, y_data, no_AbAg)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y


        else:
            x, y = self.Load_datasets_del_XXX_Datasets(x_data_available, x_data, y_data)
            random.seed(self.config.train.seed)
            random.shuffle(x)
            random.seed(self.config.train.seed)
            random.shuffle(y)
            return x, y

    @classmethod
    def sole_x_label(cls, x_list, label_list):
        tuple_list = []
        for i, j in zip(x_list, label_list):
            tuple_list.append([i, j])
        tuple_list = [i for n, i in enumerate(tuple_list) if i not in tuple_list[:n]]
        x_list = []
        label_list = []
        for i in tuple_list:
            x_list.append(i[0])
            label_list.append(i[1])
        return x_list, label_list

    def Load_total_Datasets(self, x_data_available, x_data, y_data, S1131, M1707, total):


        x_S1131_list, S1131_label_list = self.Load_S1131_Datasets(x_data_available, x_data[0], y_data, S1131, total)
        x_dataset_del_S1131_list, dataset_del_S1131_label_list = self.Load_datasets_del_XXX_Datasets(x_data_available,
                                                                                                     x_data[1], y_data,
                                                                                                     total)
        x_M1707_list, M1707_label_list = self.Load_M1707_Datasets(x_data_available, x_data[2], y_data, M1707, total)
        x_dataset_del_M1707_list, dataset_del_M1707_label_list = self.Load_datasets_del_XXX_Datasets(x_data_available,
                                                                                                     x_data[3], y_data,
                                                                                                     total)

        x_list = x_S1131_list + x_dataset_del_S1131_list + x_M1707_list + x_dataset_del_M1707_list
        label_list = S1131_label_list + dataset_del_S1131_label_list + M1707_label_list + dataset_del_M1707_label_list


        x_list, label_list = self.sole_x_label(x_list, label_list)
        y_dict = {}
        x_available_list = os.listdir(x_data_available)

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            ne = line.strip().split(';')
            li = [ne[0].strip(), None, ne[1].strip()]
            li = self.pdb_neme_toAB(li)
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[2].strip(',')
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available, label_list)
        return x, y

    def Load_datasets_del_XXX_Datasets(self, x_data_available, x_data, y_data, total=False):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r').readlines()
        for line in x_data_f:
            x_list.append(line.strip())

        x_list = list(set(x_list))
        if total: return x_list, [''] * len(x_list)

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            ne = line.strip().split(';')
            li = [ne[0].strip(), None, ne[1].strip()]
            li = self.pdb_neme_toAB(li)
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[2].strip(',')
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_M1707_Datasets(self, x_data_available, x_data, y_data, M1707, total=False):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        label_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r', encoding='UTF-8')
        for line in x_data_f:
            if line[1] == '#':
                continue
            li = line.strip().split('"')
            x_list.append(self.pdb_neme_toAB(['_'.join(li[0].split(',')[: 2]), None, li[3]]))


            if li[4].split(',')[2] == 'reverse':
                label_list.append(li[4].split(',')[2] + '_')
            else:
                label_list.append('')
            print(line)
        if total: return x_list, label_list

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            ne = line.strip().split(';')
            li = [ne[0].strip(), None, ne[1].strip()]
            li = self.pdb_neme_toAB(li)
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[2].strip(',')
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available, label_list)
        return x, y

    def Load_S1131_Datasets(self, x_data_available, x_data, y_data, S1131, total=False):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        if x_data[-5:] == '.xlsx':
            x_data_f = openpyxl.load_workbook(x_data)
            x_data_f = x_data_f.active
            pdb_num = 0
            for num, line in enumerate(x_data_f.rows):
                if line[0].value == 'pdbID':
                    continue
                if num >= 1403:
                    break
                if line[6].value != 1:
                    continue
                li = [line[0].value + '_' + line[1].value + '_' + line[2].value, '', line[3].value]
                x_list.append(self.pdb_neme_toAB(li))
                y_dict[self.pdb_neme_toAB(li)] = [None, None, line[5].value]
                pdb_num += 1
                print(pdb_num)

        elif x_data[-4:] == '.csv':
            x_data_f = open(x_data, 'r')
            for line in x_data_f:
                if line[0] == '#':
                    continue
                li = line.strip().split(',')
                x_list.append(
                    self.pdb_neme_toAB([li[0] + '_' + li[1], None, mut_format_conversion(li[4])]))
                print(line)
            x_list = list(set(x_list))
            if total: return x_list, [''] * len(x_list)

            y_data_f = open(y_data, 'r')
            for line in y_data_f:
                if line[0] == '#':
                    continue
                ne = line.strip().split(';')
                li = [ne[0].strip(), None, ne[1].strip()]
                li = self.pdb_neme_toAB(li)
                try:
                    print(y_dict[str(li)])
                    y_dict[str(li)] = [0]
                except KeyError:
                    y_dict[str(li)] = ne[2].strip(',')
                print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_S4169_Datasets(self, x_data_available, x_data, y_data, S4169, total=False):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r')
        for line in x_data_f:
            if line[0] == '#':
                continue
            if line.split(',')[0] == 'protein': continue
            li = line.strip().split(',')
            x_list.append(
                self.pdb_neme_toAB([li[0] + '_' + li[1], None, mut_format_conversion(li[4])]))
            print(line)
        x_list = list(set(x_list))
        if total: return x_list, [''] * len(x_list)

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            if line.split(',')[0] == 'protein': continue
            ne = line.strip().split(',')
            li = self.pdb_neme_toAB([ne[0] + '_' + ne[1], None, mut_format_conversion(ne[4])])
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[5].strip(',')
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_S8338_Datasets(self, x_data_available, x_data, y_data, S8338, total=False):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r')
        for line in x_data_f:
            if line[0] == '#':
                continue
            if line.split(',')[0] == 'protein': continue
            li = line.strip().split(',')
            x_list.append(
                self.pdb_neme_toAB([li[0] + '_' + li[1], None, mut_format_conversion(li[4])]))
            print(line)
        x_list = list(set(x_list))
        if total: return x_list, [''] * len(x_list)

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            if line.split(',')[0] == 'protein': continue
            ne = line.strip().split(',')
            li = self.pdb_neme_toAB([ne[0] + '_' + ne[1], None, mut_format_conversion(ne[4])])
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[5].strip(',')
            print(line)

        label_list = [''] * len(x_list) + ['reverse_'] * len(x_list)
        reverse_x_list = []
        for i in x_list:
            reverse_x_list.append(i)
        x_list = x_list + reverse_x_list
        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available, label_list=label_list)
        return x, y

    def Load_S645_Datasets(self, x_data_available, x_data, y_data, S645):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        chain_dict = {}
        for i in x_available_list:
            i = i.split('_')
            if i[0] == 'HM':
                chain_dict[i[0] + '_' + i[1] + '_' + '_'.join(i[3:])] = i[2]
            else:
                chain_dict[i[0] + '_' + '_'.join(i[2:])] = i[1]

        x_data_f = openpyxl.load_workbook(x_data)
        x_data_f = x_data_f['AB_bind_645']
        pdb_num = 0
        for num, line in enumerate(x_data_f.rows):
            if num <= 2: continue
            name = line[0].value
            mut = mut_format_conversion(line[1].value)
            chain = chain_dict[name + '_' + mut]
            ddg = str(line[3].value)
            x_list.append(name + '_' + chain + '_' + mut)
            try:
                print(y_dict[name + '_' + chain + '_' + mut])
                print(name + '_' + chain + '_' + mut)
            except:
                y_dict[name + '_' + chain + '_' + mut] = ddg
            pdb_num += 1


        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_S645_no27_Datasets(self, x_data_available, x_data, y_data, S645_no27):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        chain_dict = {}
        for i in x_available_list:
            i = i.split('_')
            if i[0] == 'HM':
                chain_dict[i[0] + '_' + i[1] + '_' + '_'.join(i[3:])] = i[2]
            else:
                chain_dict[i[0] + '_' + '_'.join(i[2:])] = i[1]

        x_data_f = openpyxl.load_workbook(x_data)
        x_data_f = x_data_f['AB_bind_645']
        pdb_num = 0
        for num, line in enumerate(x_data_f.rows):
            if num <= 2: continue
            name = line[0].value
            mut = mut_format_conversion(line[1].value)
            chain = chain_dict[name + '_' + mut]
            ddg = str(line[3].value)

            if float(ddg) == 8:
                print('lalala')
                continue
            x_list.append(name + '_' + chain + '_' + mut)
            try:
                print(y_dict[name + '_' + chain + '_' + mut])
                print(name + '_' + chain + '_' + mut)
            except:
                y_dict[name + '_' + chain + '_' + mut] = ddg
            pdb_num += 1


        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_predict_Datasets(self, x_data_available, x_data, y_data, predict):
        x_available_list = os.listdir(x_data_available)
        y_dict = {}
        for i in x_available_list:
            y_dict[i] = '0'
        x, y = self.load_data(x_available_list, x_available_list, y_dict, x_data_available)
        return x, y

    def Load_C3684_Datasets(self, x_data_available, x_data, y_data, C3684):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r').readlines()
        for line in x_data_f:
            line = line.split(';')
            if line[0] == 'Entry': continue
            x_list.append('6M0J_AE_' + line[4])

        x_list = list(set(x_list))

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            line = line.split(';')
            if line[0] == 'Entry': continue
            li = '6M0J_AE_' + line[4]

            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = line[7]
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y

    def Load_no_AbAg_Datasets(self, x_data_available, x_data, y_data, no_AbAg):
        x_available_list = os.listdir(x_data_available)
        x_list = []
        y_dict = {}

        x_data_f = open(x_data, 'r')
        for line in x_data_f:
            if line[0] == '#':
                continue
            li = line.strip().split(';')
            x_list.append(self.pdb_neme_toAB([li[0], None, li[1]]))
            print(line)
        x_list = list(set(x_list))

        y_data_f = open(y_data, 'r')
        for line in y_data_f:
            if line[0] == '#':
                continue
            ne = line.strip().split(';')
            li = [ne[0].strip(), None, ne[1].strip()]
            li = self.pdb_neme_toAB(li)
            try:
                print(y_dict[str(li)])
                y_dict[str(li)] = [0]
            except KeyError:
                y_dict[str(li)] = ne[2].strip(',')
            print(line)

        x, y = self.load_data(x_available_list, x_list, y_dict, x_data_available)
        return x, y


def get_HIV_78_data_list(x_data_available):
    pdb_chain = 'G_HL'



    x_available_list = []
    for curDir, dirs, files in os.walk(x_data_available): break
    data_list = files

    for i in data_list:
        if i == 'VRC01.pdb': continue
        data = i.split('.')

        pdb_name = 'VRC01_GHL'
        mut_chain = 'G'
        pos = data[2][1: -1]
        ori_res = data[2][0]
        mut_res = data[2][-1]

        mut = ori_res + mut_chain + pos + mut_res
        x_available_list.append(pdb_name + '_' + mut)

    return x_available_list


def get_AB_228_data_list(x_data_available, y_data):

    file_path = glob.glob(x_data_available + '/*')
    pdb_model_dict = get_AB_228_pdbmodel(file_path)

    df = pd.read_csv(y_data)

    label = df.iloc[:, 0].apply(get_label)
    pdb_name = df.iloc[:, 1].apply(lambda x: x.split('.')[0]) + '_' + df.iloc[:, 1].apply(
        lambda x: ''.join(pdb_model_dict[x.split('.')[0]].child_dict.keys()))
    mut_chain = df.iloc[:, 3]
    ori_res = df.iloc[:, 2].apply(lambda x: x[0])
    pos = df.iloc[:, 2].apply(lambda x: x[1:-1])
    mut_res = df.iloc[:, 2].apply(lambda x: x[-1])

    mut_df = pd.concat([label, ori_res, mut_chain, pos, mut_res], axis=1,
                       keys=['label', 'ori_res', 'mut_chain', 'pos', 'mut_res'])
    mut = mut_df.apply(
        lambda x: x['ori_res'] + x['mut_chain'] + x['pos'] + x['mut_res'] if x['label'] != 'reverse_' else x[
                                                                                                               'mut_res'] +
                                                                                                           x[
                                                                                                               'mut_chain'] +
                                                                                                           x['pos'] + x[
                                                                                                               'ori_res'],
        axis=1)

    x_list = list(pdb_name + '_' + mut)
    label_list = list(label)

    return x_list, label_list


def get_label(label):
    if label == 'Original':
        return ''
    elif label == 'Reverse':
        return 'reverse_'


def get_HIV_78_data_path(x_data_available):
    pdb_chain = 'G_HL'

    x_available_list = []
    for curDir, dirs, files in os.walk(x_data_available): break
    data_list = files

    for i in data_list:
        if i == 'VRC01.pdb': continue
        data = i.split('.')

        pdb_name = 'VRC01_GHL'
        mut_chain = 'G'
        pos = data[2][1: -1]
        ori_res = data[2][0]
        mut_res = data[2][-1]

        mut = ori_res + mut_chain + pos + mut_res
        x_available_list.append(pdb_name + '_' + mut)

    file_path = glob.glob(x_data_available + '/*')
    x_data_path_dict = {}
    for i, j in zip(file_path, x_available_list):
        tem = i.split('/')[-1]
        if tem == 'VRC01.pdb':
            wt_path = i
            file_path.remove(i)
            break
    for i, j in zip(file_path, x_available_list):
        tem = i.split('/')[-1]
        if tem == 'VRC01.pdb':
            continue
        assert tem.split('.')[-2][1:-1] == j.split('_')[-1][2:-1]
        mut_path = i
        x_data_path_dict[j] = [mut_path, wt_path]

    return x_data_path_dict


def get_AB_228_pdbmodel(file_path):
    out_dict = {}
    warnings.simplefilter('ignore', BiopythonWarning)
    parser = PDBParser()
    for f in file_path:
        if len(f.rsplit('/')[-1].split('.')) == 2:
            out_dict[f.rsplit('/')[-1].split('.')[0]] = parser.get_structure(None, f)[0]
    return out_dict


def get_AB_228_data_path(x_data_available, y_data):
    df = pd.read_csv(y_data)

    x_available_list = []
    for curDir, dirs, files in os.walk(x_data_available): break
    data_list = files

    file_path = glob.glob(x_data_available + '/*')
    pdb_model_dict = get_AB_228_pdbmodel(file_path)

    for i in data_list:
        if len(i.rsplit('/')[-1].split('.')) == 2: continue
        data = i.split('.')

        label = ''
        pdb_name = data[0] + '_' + ''.join(pdb_model_dict[data[0]].child_dict.keys())
        mut_chain =\
        df[(df['PDB'].apply(lambda x: x.split('.')[0]) == data[0]) & (df['Mutation'] == data[-2])]['Chain'].iloc[0]
        pos = data[-2][1: -1]
        ori_res = data[-2][0]
        mut_res = data[-2][-1]

        mut = ori_res + mut_chain + pos + mut_res
        x_available_list.append(label + pdb_name + '_' + mut)

    x_data_path_dict = {}
    WT_path_dict = {}
    for i, j in zip(file_path, x_available_list):
        tem = i.split('/')[-1]
        if len(tem.rsplit('/')[-1].split('.')) == 2:
            WT_path_dict[tem.split('.')[0]] = i
            file_path.remove(i)

    for i, j in zip(file_path, x_available_list):
        tem = i.split('/')[-1]
        assert tem.split('.')[-2][1:-1] == j.split('_')[-1][2:-1]
        mut_path = i
        x_data_path_dict[j] = [mut_path, WT_path_dict[tem.split('.')[0]]]

    return x_data_path_dict


def get_HIV_78_y(y_data):
    df = pd.read_csv(y_data)
    out = get_df_dict(df)
    return out


def get_AB_228_y(y_data, x_list, label_list):
    x_list = [l + x for x, l in zip(x_list, label_list)]
    out_dict = {}
    df = pd.read_csv(y_data)
    for k in x_list:
        if k.split('_')[0] == 'reverse':
            value = list(df[(df['Mutation type'] == 'Reverse') &
                            (df['PDB'].apply(lambda x: x.split('.')[0]) == k.split('_')[1]) &
                            (df['Mutation'] == (k.split('_')[-1][-1] + k.split('_')[-1][2: -1] + k.split('_')[-1][0]))][
                             'Experimental'])

        else:
            value = list(df[(df['Mutation type'] == 'Original') &
                            (df['PDB'].apply(lambda x: x.split('.')[0]) == k.split('_')[0]) &
                            (df['Mutation'] == (k.split('_')[-1][0] + k.split('_')[-1][2:]))]['Experimental'])

        out_dict[k] = y_reverse(str(statistics.mean(value)))

    return out_dict


def y_reverse(str_y):
    if str_y[0] == '-':
        str_y = str_y[1:]
    else:
        str_y = '-' + str_y
    return str_y


def get_df_dict(df):
    df_filtered = df[df.iloc[:, 0] == 'ORIGINAL']


    out_keys = 'VRC01_GHL' + '_' + df_filtered.iloc[:, 2].apply(get_df_ori_res) + 'G' + df_filtered.iloc[:, 2].apply(
        get_df_posmut_res)









    out_values = df_filtered.loc[:, 'Exptal'].apply(cal_HIV_ddg_4)

    out = dict(zip(out_keys, out_values))
    return out


def cal_HIV_ddg_1(x):
    if x == 0: return str(8)
    kd_wt = 5.623E-10
    kd_mut = kd_wt / (x / 100)
    T = 273.15 + 25
    ddg = str((8.314 / 4184) * T * math.log(float(kd_mut) / float(kd_wt), math.e))
    return str(ddg)


def cal_HIV_ddg_2(x):
    if x == 0: return str(8)
    T = 273.15 + 25
    ddg = str((8.314 / 4184) * T * math.log(float(x / 100), math.e))
    return str(ddg)


def cal_HIV_ddg_3(x):
    if x == 0: return str(8)
    kd_wt = 5.623E-10
    kd_mut = kd_wt * (x / 100)
    T = 273.15 + 25
    ddg = str((8.314 / 4184) * T * math.log(float(kd_mut) / float(kd_wt), math.e))
    return str(ddg)


def cal_HIV_ddg_4(x):
    if x == 0: return str(8)
    kd_wt = 5.623E-10
    kd_mut = kd_wt * (x / 100)
    T = 273.15 + 25
    ddg = str((8.314 / 4184) * T * math.log(x / 100, math.e))
    return str(ddg)


def get_df_ori_res(column):
    return column[-1]


def get_df_posmut_res(column):
    return column[1:-1] + column[0]


def get_protein_divide(pdb_name):

    chains = []
    tem_data = open('../datasets/skempi_v2.csv').readlines()
    for line in tem_data:
        if line[0] == '#': continue
        name = line.split(';', 1)[0]
        if CustomDataset.pdb_neme_toAB(name) == pdb_name.split('_')[0] + '_' + pdb_name.split('_')[1]:
            chains = name.split('_', 1)[1].split('_')
    return chains


def batch_contact_area(config, x_batch, contact_area_dict):
    out_list = []
    x_name = x_batch['pdb_name']
    for i in x_name:
        if i.split('_')[0] == 'reverse':
            out_list.append('-' + contact_area_dict[i.split('_', 1)[1]])
        out_list.append(contact_area_dict[i])
    out_list = y_recursive_to(out_list, config.model.device)
    x_batch['contact_area'] = out_list
    return x_batch


def mut_format_conversion(muts):
    muts_ori = muts
    muts = muts.split(';')
    if len(muts) == 1:
        muts = muts[0].split(', ')
    out = []
    for mut in muts:
        if (':' not in mut) and ('_' not in mut):
            return muts_ori
        mut = mut.split(':')
        if len(mut) == 1:
            mut = mut[0].split('_')
        chain = mut[0]
        original_res = mut[1][0]
        pos = mut[1][1: -1]
        after_res = mut[1][-1]
        mut = original_res + chain + pos + after_res
        out.append(mut)
    return ','.join(out)


def save_variable(v, filename):
    f = open(filename, 'wb')
    pickle.dump(v, f)
    f.close()
    return filename


def load_variable(filename):
    f = open(filename, 'rb')
    r = pickle.load(f)
    f.close()
    return r


def train_val_fold(path, fold=5):
    data = {}
    tem_v = []
    for n in path:
        tem_v.append(path[n])
        data[n] = {'test': tem_v}
        tem_v = []

        tem_t = []
        for m in path:
            if n == m:
                continue
            tem_t.append(path[m])
        data[n]['train'] = tem_t
    return data


def load_data(config, fold):
    val = load_variable(fold['test'][0])
    train_x = []
    train_y = []
    for train_path in fold['train']:
        data = load_variable(train_path)
        for x, y in data:
            train_x.append(x)
            train_y.append(y)
    train = CustomDataset(config, x_data=train_x, y_data=train_y)
    return train, val


def data_shuffle(config, com_data, seed):
    x = []
    y = []
    for n, m in com_data:
        x.append(n)
        y.append(m)
    random.seed(seed)
    random.shuffle(x)
    random.seed(seed)
    random.shuffle(y)
    data = CustomDataset(config, x_data=x, y_data=y)
    return data


def load_weight(weight):
    import collections
    tem_OrderedDict = collections.OrderedDict()
    for k, v in weight.items():
        if k.split('.', 1)[0] != '_orig_mod': return weight
        tem_k = k.split('.', 1)[1]
        tem_OrderedDict[tem_k] = v
    return tem_OrderedDict


def get_file_existence_time(file):
    timestamp = os.path.getctime(file)
    file_create_time = datetime.datetime.fromtimestamp(timestamp)
    current_time = datetime.datetime.now()
    delta = current_time - file_create_time
    try:
        existence_min = str(delta).split(',')[0].split(' ')[0] + '* 24 * 60 + ' +\
                        str(int(str(delta).split(',')[1].split(':')[0])) + ' * 60 + ' +\
                        str(int(str(delta).split(',')[1].split(':')[1]))
    except IndexError:
        existence_min = str(int(str(delta).split(':')[0])) + ' * 60 + ' +\
                        str(int(str(delta).split(':')[1]))

    loc = {}
    exec('existence_hour = (' + existence_min + ')/ float(60)', globals(), loc)
    existence_hour = loc['existence_hour']
    return existence_hour


def del_procedure_parameter(files_path, num=100):
    f_list = []
    for _, _, f_list in os.walk(files_path): break

    f_path_list = []
    f_existence_time = []
    for f in f_list:
        f_path = files_path + '/' + f
        f_path_list.append(f_path)
        f_existence_time.append(get_file_existence_time(f_path))


    while len(f_path_list) > num:
        max_index = f_existence_time.index(max(f_existence_time))
        os.remove(f_path_list[max_index])
        del f_path_list[max_index]
        del f_existence_time[max_index]
    return


def save_procedure_parameter(config, validation_flag, label, epoch, py_name,
                             train_Rp, train_losses_epoch,
                             validation_Rp, validation_losses_epoch,
                             test_Rp, test_losses_epoch):
    if validation_flag:
        f_path = config.observer.parameter.parameter_path + '/' + label + '/validation_curve/parameter_' +\
                 py_name.split('.')[0] + '/' + str(epoch) + '_' + str(
            datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')) + '.txt'
    else:
        f_path = config.observer.parameter.parameter_path + '/' + label + '/parameter_' + py_name.split('.')[
            0] + '/' + str(epoch) + '_' + str(
            datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')) + '.txt'

    os.makedirs(os.path.dirname(f_path), exist_ok=True)

    procedure_parameter_f = open(f_path, 'w')

    procedure_parameter_f.write('train_Rp' + '\n')
    procedure_parameter_f.write(str(train_Rp) + '\n')

    procedure_parameter_f.write('train_losses_epoch' + '\n')
    procedure_parameter_f.write(str(train_losses_epoch) + '\n')

    procedure_parameter_f.write('validation_Rp' + '\n')
    procedure_parameter_f.write(str(validation_Rp) + '\n')

    procedure_parameter_f.write('validation_losses_epoch' + '\n')
    procedure_parameter_f.write(str(validation_losses_epoch) + '\n')

    procedure_parameter_f.write('test_Rp' + '\n')
    procedure_parameter_f.write(str(test_Rp) + '\n')

    procedure_parameter_f.write('test_losses_epoch' + '\n')
    procedure_parameter_f.write(str(test_losses_epoch) + '\n')

    procedure_parameter_f.close()


    del_procedure_parameter(os.path.dirname(f_path), num=100)


def load_terminal_paras(config, load_result_weight=False, validation_flag=False, label=None):
    if load_result_weight is False:








        validation_flag = False
        label = 'S1131_DG_random'
    else:
        validation_flag = validation_flag
        label = label
        print('load_result_weight!!!!!!')


    config.datasets.label = label
    if validation_flag:
        config.observer.validation_Rp = True
        config.observer.validation_losses_epoch = True
    else:
        config.observer.validation_Rp = False
        config.observer.validation_losses_epoch = False
    return config


def str_to_bool(x):
    print(x)
    if x == 'True':
        return True
    elif x == 'False':
        return False


def get_compound_class_data(skempi_format_data_f):
    compound = {}
    compound_name = []

    len_skempi_format_data = 0
    for line in skempi_format_data_f:
        if line[0] == '#':
            header = line
            continue
        li = line.strip().split(';')
        compound_name.append(li[0])
        len_skempi_format_data += 1
        try:
            compound[li[0]].append(line)
        except KeyError:
            compound[li[0]] = [line]

    return compound, compound_name, len_skempi_format_data


def get_header(skempi_format_data_f):
    head = ''
    for row in skempi_format_data_f:
        if row[0] == '#':
            head = row
            break
    return head


def introduce_contact_area(config, residue_data, contact_area):
    if config.feature.contact_area[1] == 'Residue level':
        N, L, dim = residue_data.size()
        contact_area = contact_area[:, None, None].expand(-1, L, -1)
        residue_data = torch.cat((residue_data, contact_area), dim=-1)
    elif config.feature.contact_area[1] == 'Compound level (MLP)':
        pass
    elif config.feature.contact_area[1] == 'Compound level (Manual design)':
        pass

    return residue_data


def update_requires_grad(config, model, bool_=False):
    if bool_:



        optimizer = optim.Adam(
            filter(lambda p: p.requires_grad, model.parameters()),
            lr=config.train.optimizer.lr,
            betas=(config.train.optimizer.beta1, config.train.optimizer.beta2),
            weight_decay=config.train.optimizer.weight_decay)
    else:
        optimizer = optim.Adam(
            model.parameters(),
            lr=config.train.optimizer.lr,
            betas=(config.train.optimizer.beta1, config.train.optimizer.beta2),
            weight_decay=config.train.optimizer.weight_decay)
    return optimizer


def freeze_model(model, to_freeze_dict, keep_step=None):
    for (name, param) in model.named_parameters():
        if name in to_freeze_dict:
            param.requires_grad = False
    return model


def get_freeze_weight(weight):
    import collections
    tem_OrderedDict = collections.OrderedDict()
    for k, v in weight.items():
        if k.split('.', 1)[0] == 'ddG_readout': continue




        tem_OrderedDict[k] = v
    return tem_OrderedDict


def loadIGMIfeature(config, case):
    root_path = Path(config.datasets_DG.readyPDB)
    paths_pdb = case['name']
    proteinA = case['proteinA'].replace(",", "").replace(" ", "")
    proteinB = case['proteinB'].replace(",", "").replace(" ", "")
    proteinA = [it for it in proteinA]
    proteinB = [it for it in proteinB]
    proteinA = "".join(proteinA)
    proteinB = "".join(proteinB)
    set_name = case['set']
    pdb_name = [paths_pdb + '_' + proteinA + '_' + proteinB, set_name]

    mutations = case['Mutations']
    if not pd.isna(mutations):
        mutations = mut_format_conversion(mutations)
        paths_pdb = paths_pdb + '_' + ''.join(proteinA) + ''.join(proteinB) + '_' + '_'.join(mutations.split(','))

    pdb_sp = paths_pdb.split('_')
    if set_name == "SKEMPI_v2.0":
        data_root = Path(config.datasets_DG.Skempi2)
        pdb_sp = paths_pdb.split('_')
        if len(pdb_sp) > 2:
            dir = data_root / paths_pdb
            path = dir / ('MUT_' + paths_pdb + ".pdb")
            energy_path = dir / "mut_energy.sc"
            h5_path = dir / ('MUT_' + paths_pdb + ".h5")
            region_path = dir / ('MUT_' + paths_pdb + ".region")
        else:
            file = list(data_root.glob(f"{paths_pdb}*"))[0]
            name = str(file.name)
            dir = data_root / file
            path = dir / ('WT_' + name + ".pdb")
            energy_path = dir / "wt_energy.sc"
            h5_path = dir / ('WT_' + name + ".h5")
            region_path = dir / ('WT_' + name + ".region")
    else:
        dir = root_path / set_name / (paths_pdb + '_' + proteinA + '_' + proteinB)
        path = dir / (paths_pdb + '_' + proteinA + proteinB + ".pdb")
        energy_path = dir / "wt_energy.sc"
        region_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".region")
        h5_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".h5")




    data_dg = parse_pdb(path, pdb_name)

    _, interface_mask = get_interface_residues(config, data_dg, pdb_name)

    data_dg['interface_label'] = interface_mask.int()
    if len(pdb_sp) > 2:
        data_dg['pdb_name'] = pdb_name[0] + '_' + '_'.join(mutations.split(',')) + '&' + pdb_name[1]
    else:
        data_dg['pdb_name'] = '&'.join(pdb_name)

    return data_dg


def Aggregate_features(x0, x1):
    out_dict = x1

    out_dict['res_prottrans'] = torch.tensor(x0['res_prottrans']).to(out_dict['pos14'])
    out_dict['res_energy'] = torch.tensor(x0['res_energy']).to(out_dict['pos14'])
    out_dict['neighbors'] = torch.tensor(x0['neighbors']).to(out_dict['pos14'])
    out_dict['interface_label'] = torch.tensor(x0['interface_label']).to(out_dict['pos14'])
    return out_dict


def get_contact_dis_mat(pos_a, pos_b, mask_a, mask_b):
    l1, n_atom = pos_a.shape[:2]
    l2, n_atom = pos_b.shape[:2]

    mask_a = mask_a.astype('bool')
    mask_b = mask_b.astype('bool')
    p1 = pos_a[mask_a]
    p2 = pos_b[mask_b]
    dis_mat = cdist(p1, p2)


    index_a = np.tile(np.arange(l1)[:, None], reps=[1, n_atom])[mask_a]
    index_b = np.tile(np.arange(l2)[:, None], reps=[1, n_atom])[mask_b]


    contact_dis_mat = dis_mat

    groups = index_a
    _ndx = np.argsort(groups)
    _id, _pos, g_count = np.unique(groups[_ndx], return_index=True, return_counts=True)
    contact_dis_mat = np.minimum.reduceat(contact_dis_mat[_ndx], _pos, axis=0)

    groups = index_b
    _ndx = np.argsort(groups)
    _id, _pos, g_count = np.unique(groups[_ndx], return_index=True, return_counts=True)
    contact_dis_mat = np.minimum.reduceat(contact_dis_mat[:, _ndx], _pos, axis=1)

    return contact_dis_mat


def single_mutation_parsing(mutation):
    aatype_from = mutation[0]
    chainid = mutation[1]
    position = int(mutation[2:-1])

    aatype_to = mutation[-1]
    return chainid, position, aatype_from, aatype_to


def mutation_parsing_to_frame(mutation):
    cols = 'chainid, position, aatype_from, aatype_to'.split(', ')
    mutations = []
    for m in mutation.split(';'):
        mutations += [single_mutation_parsing(m)]
    df_mut = pd.DataFrame(mutations, columns=cols)
    df_mut['position'] = df_mut['position'].astype('int')
    return df_mut


def get_contact_node(adjacency_matrix, max_len, limit_length):

    edges = []
    num_nodes = len(adjacency_matrix)

    for i in range(num_nodes):
        for j in range(i + 1, num_nodes):
            if adjacency_matrix[i, j] > 0:
                edges.append((i, j, adjacency_matrix[i, j]))


    edges.sort(key=lambda x: x[2])


    selected_nodes = set()
    for edge in edges:
        u, v, weight = edge
        if weight == 10000:
            break
        if weight > limit_length:
            break
        selected_nodes.add(u)


        selected_nodes.add(v)




    return selected_nodes


def get_knn_dis_index(pos37, mask37, chainids, limit_length, n_neighbor, max_len=6000):
    pos37 = pos37[:max_len]
    mask37 = mask37[:max_len]
    chainid = chainids[:max_len]



    ca_wt = pos37[:, 1]


    contact_dis_mat = pairwise_distances(ca_wt, ca_wt)


    contact_dis_mat[chainid[None, :] == chainid[:, None]] = 10000






    inter_res_index = get_contact_node(contact_dis_mat, n_neighbor, limit_length)









    return list(inter_res_index)


def get_knn_dis_index_mat(pos14, n_neighbor=None):
    if n_neighbor == None:
        return np.arange(pos14.shape[0])
    ca_pos = pos14[:, ATOM_CA]

    dis_mat = pairwise_distances(ca_pos, ca_pos)
    dis_mat_zero_mask = dis_mat == 0
    dis_mat[dis_mat_zero_mask] = np.inf
    dis_min_arg = dis_mat.argsort(axis=-1)
    inter_res_index = dis_min_arg[:, :n_neighbor]

    pad_cols = n_neighbor - inter_res_index.shape[1]
    inter_res_index = np.pad(inter_res_index, ((0, 0), (0, pad_cols)), mode='constant')
    return inter_res_index




def energy_filter(energyfile):
    del_labels = ['SCORE:', 'pose_id', "fa_intra_rep", 'fa_intra_sol_xover4', 'lk_ball_wtd', 'pro_close',
                  'hbond_sr_bb', 'dslf_fa13', 'omega', 'fa_dun', 'p_aa_pp', 'yhh_planarity',
                  'ref', 'rama_prepro', 'total', 'description']
    energyfile = energyfile[energyfile['restype2'] != 'onebody']
    energyfile = energyfile.drop(del_labels, axis=1)
    energyfile[['resi1', 'resi2']] = energyfile[['resi1', 'resi2']].astype(int)

    len_res = max(energyfile['resi2'])
    energy_matrix = np.zeros((len_res, len_res, 7))
    for _, row in energyfile.iterrows():
        energy_matrix[row['resi1'] - 1][row['resi2'] - 1] = row['fa_atr':]
        energy_matrix[row['resi2'] - 1][row['resi1'] - 1] = row['fa_atr':]
    return energy_matrix




class DGComplexLoader(object):
    def __init__(self, data_root):
        super(DGComplexLoader, self).__init__()
        self.data_root = Path(data_root)


    def get_protein(self, paths_pdb, proteinA, proteinB, set):



































        proteinA = "".join(proteinA)
        proteinB = "".join(proteinB)
        if set == "SKEMPI_v2.0":

            data_root = Path(config.datasets_DG.Skempi2)


            pdb_sp = paths_pdb.split('_')
            if len(pdb_sp) > 2:
                dir = data_root / paths_pdb
                path = dir / ('MUT_' + paths_pdb + ".pdb")
                energy_path = dir / "mut_energy.sc"
                h5_path = dir / ('MUT_' + paths_pdb + ".h5")
                region_path = dir / ('MUT_' + paths_pdb + ".region")
                if not os.path.exists(path):
                    print(path)
                    return None, None, None, None
            else:
                file = list(data_root.glob(f"{paths_pdb}*"))[0]
                name = str(file.name)
                dir = data_root / file
                path = dir / ('WT_' + name + ".pdb")
                energy_path = dir / "wt_energy.sc"
                h5_path = dir / ('WT_' + name + ".h5")
                region_path = dir / ('WT_' + name + ".region")
        else:
            dir = self.data_root / set / (paths_pdb + '_' + proteinA + '_' + proteinB)
            path = dir / (paths_pdb + '_' + proteinA + proteinB + ".pdb")
            energy_path = dir / "wt_energy.sc"
            region_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".region")
            h5_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".h5")
        pdb_string = path.read_text()


        pdbs = proteins2.ProteinInput.from_pdb(io.StringIO(pdb_string),
                                               with_angles=True, return_dict=True)


        energyfile = pd.read_csv(energy_path, delim_whitespace=True, header=0)
        energy_matrix = energy_filter(energyfile)

        region = pd.read_csv(region_path)
        prottrans = h5py.File(h5_path, 'r')
        return pdbs, energy_matrix, region, prottrans


    def get_chains(self, paths_pdb, proteinA, proteinB, set):
        proteinA = "".join(proteinA)
        proteinB = "".join(proteinB)


        if set == "SKEMPI_v2.0":
            data_root = Path(config.datasets_DG.Skempi2)


            pdb_sp = paths_pdb.split('_')
            if len(pdb_sp) > 2:
                dir = data_root / paths_pdb
                path = dir / ('MUT_' + paths_pdb + ".pdb")
                energy_path = dir / "mut_energy.sc"
                h5_path = dir / ('MUT_' + paths_pdb + ".h5")
                region_path = dir / ('MUT_' + paths_pdb + ".region")
            else:
                file = list(data_root.glob(f"{paths_pdb}*"))[0]
                name = str(file.name)
                dir = data_root / file
                path = dir / ('WT_' + name + ".pdb")
                energy_path = dir / "wt_energy.sc"
                h5_path = dir / ('WT_' + name + ".h5")
                region_path = dir / ('WT_' + name + ".region")
        else:
            dir = self.data_root / set / (paths_pdb + '_' + proteinA + '_' + proteinB)
            path = dir / (paths_pdb + '_' + proteinA + proteinB + ".pdb")
            energy_path = dir / "wt_energy.sc"
            region_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".region")
            h5_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".h5")






        result = []

        parser = PDB.PDBParser()


        structure = parser.get_structure('0', path)
        model = structure[0]


        for chain in model:
            result.append(chain.id)
        return result

    def load(self, paths_pdb, proteinA, proteinB, set) -> (ProteinInput):
        if isinstance(paths_pdb, str):
            paths_wt = paths_pdb.split(',')

        pdbs, energy, region, prottrans = self.get_protein(paths_pdb, proteinA, proteinB, set)

        index_cor = region[region["region"] == "COR"]['pos'].tolist()
        index_rim = region[region["region"] == "RIM"]['pos'].tolist()
        index_sup = region[region["region"] == "SUP"]['pos'].tolist()
        conbine = index_cor + index_sup + index_rim
        conbine = conbine[:128]
        inter_region_index = [x - 1 for x in conbine]

        chains = self.get_chains(paths_pdb, proteinA, proteinB, set)
        if ("XXXX|" + chains[0]) in prottrans.keys():
            res_prottrans = np.concatenate([prottrans['XXXX|' + x][:] for x in chains], axis=0)
        elif ("PDB|" + chains[0]) in prottrans.keys():
            res_prottrans = np.concatenate([prottrans['PDB|' + x][:] for x in chains], axis=0)

        pdb_all: proteins2.ProteinInput = proteins2.proteins_merge([pdbs[chain] for chain in chains], chains)

        mask = np.zeros(pdb_all.length).astype('bool')
        mask[inter_region_index] = True

        p14 = pdb_all.to_atom14()
        interface_label = mask
        return p14, energy, res_prottrans, interface_label




class ProteinDataset(Dataset):
    def __init__(self, df_path, data_root, max_length=256,
                 col_pdb='name', col_proteinA='proteinA', col_proteinB='proteinB', col_allchains='all_chains',
                 set_c='set',
                 mutations='Mutations',
                 cols_label=None,
                 n_neighbors=None, limit_length=None,
                 train=True, diskcache=None, ):

        if cols_label is None:
            cols_label = ['dG']

        self.data_loader = DGComplexLoader(data_root)

        if isinstance(df_path, pd.DataFrame):
            self.df = df_path.copy()
        else:
            self.df = pd.read_csv(df_path, low_memory=False)





        self.max_length = max_length
        self.n_neighbors = n_neighbors
        self.limit_length = limit_length
        self.col_pdb = col_pdb
        self.col_proteinA = col_proteinA
        self.col_proteinB = col_proteinB
        self.col_allchains = col_allchains
        self.cols_label = cols_label
        self.set_c = set_c
        self.mutations = mutations
        self.train = train
        self.diskcache = diskcache


        for col in self.cols_label:
            self.df[col] = self.df.get(col, np.nan)
        self.num_classes = len(self.cols_label)

    def __len__(self):
        return len(self.df)


    def load_data(self, paths_pdb, proteinA, proteinB, cache, set):
        key = paths_pdb, proteinA, proteinB, set

        if cache is None:
            values = self.data_loader.load(*key)
            return values

        if key in cache:
            values = cache[key]
            return values

        values = self.data_loader.load(*key)
        cache[key] = values
        return values


    def get_data(self, idx_case):

        case = idx_case
        paths_pdb = case[self.col_pdb]
        proteinA = case[self.col_proteinA].replace(",", "").replace(" ", "")
        proteinB = case[self.col_proteinB].replace(",", "").replace(" ", "")
        proteinA = [it for it in proteinA]
        proteinB = [it for it in proteinB]

        set_c = case[self.set_c]
        cache = self.diskcache

        mutations = case[self.mutations]
        if not pd.isna(mutations):
            mutations = mut_format_conversion(mutations)
            paths_pdb = paths_pdb + '_' + ''.join(proteinA) + ''.join(proteinB) + '_' + '_'.join(mutations.split(','))

        values = self.load_data(paths_pdb, proteinA, proteinB, cache, set_c)
        p14 = values[0]
        energy_wt = values[1]
        res_prottrans = values[2]
        interface_label = values[3]
        labels = case[self.cols_label].fillna(0).astype('float32').values
        labels_valid_mask = case[self.cols_label].notna().values

        complex_all = {
            'pos14': p14.atom_positions.astype('float32'),
            'pos14_mask': p14.atom_mask.astype('bool'),
            'aa': p14.aatype.astype('int'),
            'seq': p14.residue_index.astype('int'),
            'chain_id': p14.chainid.astype('str'),
            'chain_seq': p14.chainseq.astype('int'),
            'res_energy': energy_wt.astype('float32'),



            'res_prottrans': res_prottrans.astype('float32'),
            'interface_label': interface_label.astype('bool')
        }









        complex_all['neighbors'] = get_knn_dis_index_mat(complex_all['pos14'],
                                                         self.n_neighbors)


        label = (labels.astype('float32'), labels_valid_mask.astype('bool'), set_c)

        return (complex_all,), label

    def __getitem__(self, idx):
        try:
            return self.get_data(idx)
        except Exception as e:
            case = self.df.iloc[idx]
            paths_wt = case[self.col_pdb]
            print(case, paths_wt)
            print(e)
            traceback.print_exc()
            raise e


def load_PredictDG_datasets(config, data_index_list, data_df_list):
    df = pd.read_excel(config.datasets_DG.DGdatasets)

    data = ProteinDataset(df, config.datasets_DG.readyPDB, cols_label=['dG'],
                          max_length=config.feature.DG_surface_num,
                          n_neighbors=64, limit_length=10, train=True,
                          diskcache=None)
    x_list = []
    y_list = []
    for j in data_df_list.iterrows():
        j = j[1]

        if pdb_exist(j):
            with open('../error/' + j['name'] + ';' + j['Mutations'], 'w') as f: pass
            continue

        x0 = data[j]
        yddg = x0[1]
        x0 = x0[0][0]
        x1 = loadIGMIfeature(config, j)
        x = Aggregate_features(x0, x1)

        batch = choice_residue(config, x, pdb_name=x['pdb_name'].split('&')[0])

        x_list.append(batch)
        y_list.append(str(yddg[0][0]))

    data = CustomDataset(config, x_data=x_list, y_data=y_list)
    return data


def pdb_exist(case):
    root_path = Path(config.datasets_DG.readyPDB)
    paths_pdb = case['name']
    proteinA = case['proteinA'].replace(",", "").replace(" ", "")
    proteinB = case['proteinB'].replace(",", "").replace(" ", "")
    proteinA = [it for it in proteinA]
    proteinB = [it for it in proteinB]
    proteinA = "".join(proteinA)
    proteinB = "".join(proteinB)
    set_name = case['set']

    mutations = case['Mutations']
    if not pd.isna(mutations):
        mutations = mut_format_conversion(mutations)
        paths_pdb = paths_pdb + '_' + ''.join(proteinA) + ''.join(proteinB) + '_' + '_'.join(mutations.split(','))

    if set_name == "SKEMPI_v2.0":

        data_root = Path(config.datasets_DG.Skempi2)


        pdb_sp = paths_pdb.split('_')
        if len(pdb_sp) > 2:
            dir = data_root / paths_pdb
            path = dir / ('MUT_' + paths_pdb + ".pdb")
            energy_path = dir / "mut_energy.sc"
            h5_path = dir / ('MUT_' + paths_pdb + ".h5")
            region_path = dir / ('MUT_' + paths_pdb + ".region")
        else:
            file = list(data_root.glob(f"{paths_pdb}*"))[0]
            name = str(file.name)
            dir = data_root / file
            path = dir / ('WT_' + name + ".pdb")
            energy_path = dir / "wt_energy.sc"
            h5_path = dir / ('WT_' + name + ".h5")
            region_path = dir / ('WT_' + name + ".region")
    else:
        dir = root_path / set_name / (paths_pdb + '_' + proteinA + '_' + proteinB)
        path = dir / (paths_pdb + '_' + proteinA + proteinB + ".pdb")
        energy_path = dir / "wt_energy.sc"
        region_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".region")
        h5_path = dir / (paths_pdb + '_' + proteinA + proteinB + ".h5")



    flag = os.path.exists(path)
    if flag:
        return False
    else:
        return True
